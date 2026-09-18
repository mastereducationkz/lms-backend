"""The register as the spreadsheet head teachers already know.

One tab per programme, teachers down, days across, «Пропуски» and «Наказание» under each date —
the shape of «Attendance and Late lessons 16.09-31.10», so payroll opens it and recognises it.
"""
from __future__ import annotations

from datetime import date
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.services.excel_export_service import sanitize_spreadsheet_value

_HEAD = Font(bold=True)
_TITLE_FILL = PatternFill("solid", fgColor="EFF3F8")
_LATE_FILL = PatternFill("solid", fgColor="FFF4E5")
_MISS_FILL = PatternFill("solid", fgColor="FDE8E8")
_CENTRE = Alignment(horizontal="center")

TOTAL_HEADERS = ("Опозданий, мин", "Пропусков", "Штраф, ₸")


def _day_label(iso: str) -> str:
    return date.fromisoformat(iso).strftime("%d.%m.%Y")


def _what_happened(cell: dict) -> Optional[str]:
    """The «Пропуски» column in words: what the LMS saw that day."""
    if not cell:
        return None
    parts = []
    if cell.get("misses"):
        parts.append("Пропуск" if cell["misses"] == 1 else f"Пропусков: {cell['misses']}")
    if cell.get("late_minutes"):
        parts.append(f"{cell['late_minutes']} мин")
    if cell.get("early_minutes"):
        parts.append(f"−{cell['early_minutes']} мин")
    if not parts and cell.get("unmeasurable"):
        return "Нет данных"
    return ", ".join(parts) or None


def workbook_for(register: dict) -> Workbook:
    """The whole period, one sheet per programme."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    days = register["days"]

    programs: dict[str, list[dict]] = {}
    for teacher in register["teachers"]:
        programs.setdefault(teacher.get("program") or "—", []).append(teacher)

    for program, teachers in programs.items():
        sheet = workbook.create_sheet(title=(program or "—")[:31])
        sheet["A1"] = "ФИО"
        sheet["A1"].font = _HEAD
        sheet["A2"] = register["period"]["label"]
        sheet["A2"].font = _HEAD
        sheet["A2"].fill = _TITLE_FILL
        sheet.column_dimensions["A"].width = 32

        for index, day in enumerate(days):
            first = 2 + index * 2
            head = sheet.cell(row=1, column=first, value=_day_label(day))
            head.font = _HEAD
            head.alignment = _CENTRE
            sheet.merge_cells(start_row=1, start_column=first, end_row=1, end_column=first + 1)
            sheet.cell(row=2, column=first, value="Пропуски").font = _HEAD
            sheet.cell(row=2, column=first + 1, value="Наказание").font = _HEAD
            for column in (first, first + 1):
                sheet.column_dimensions[get_column_letter(column)].width = 13

        totals_at = 2 + len(days) * 2
        for offset, header in enumerate(TOTAL_HEADERS):
            cell = sheet.cell(row=1, column=totals_at + offset, value=header)
            cell.font = _HEAD
            sheet.column_dimensions[get_column_letter(totals_at + offset)].width = 15

        for row_index, teacher in enumerate(teachers, start=3):
            sheet.cell(row=row_index, column=1, value=sanitize_spreadsheet_value(teacher["name"]))
            for index, day in enumerate(days):
                cell = (teacher.get("days") or {}).get(day) or {}
                what = _what_happened(cell)
                if what is None and not cell.get("fine"):
                    continue
                first = 2 + index * 2
                happened = sheet.cell(row=row_index, column=first, value=what)
                happened.alignment = _CENTRE
                if cell.get("misses"):
                    happened.fill = _MISS_FILL
                elif cell.get("late_minutes") or cell.get("early_minutes"):
                    happened.fill = _LATE_FILL
                if cell.get("fine"):
                    money = sheet.cell(row=row_index, column=first + 1, value=int(cell["fine"]))
                    money.alignment = _CENTRE

            totals = teacher.get("totals") or {}
            for offset, key in enumerate(("late_minutes", "misses", "fine")):
                sheet.cell(row=row_index, column=totals_at + offset, value=int(totals.get(key) or 0))

        foot = 3 + len(teachers)
        sheet.cell(row=foot, column=1, value="Итого").font = _HEAD
        for offset, key in enumerate(("late_minutes", "misses", "fine")):
            cell = sheet.cell(row=foot, column=totals_at + offset,
                              value=int((register["totals"] or {}).get(key) or 0))
            cell.font = _HEAD

        sheet.freeze_panes = "B3"

    if not programs:  # an empty period still opens, with its heading
        sheet = workbook.create_sheet(title="—")
        sheet["A1"] = "ФИО"
        sheet["A2"] = register["period"]["label"]
    return workbook
