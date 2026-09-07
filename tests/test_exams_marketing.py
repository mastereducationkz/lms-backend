"""Marketing eligibility of exam results: a score above the threshold OR a testimonial.

The sales team may feature a student whose current SAT attempt is above 1400 even when
no testimonial was ever collected. These tests pin the threshold (strictly greater
than), the "current attempt only" rule, which exam types the threshold applies to, and
that the two grounds are reported separately - a bare score carries no consent record,
so the sales team must be able to tell it from a consented testimonial.
"""
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import load_workbook

# The shim must be imported before any domain model module: importing a domain model
# first re-enters the partially-initialized src.models package and raises ImportError.
from src.schemas.models import ExamResult  # noqa: F401  (import-order guard)
from src.exams.marketing import (
    BASIS_SCORE,
    BASIS_TESTIMONIAL,
    MARKETING_SCORE_THRESHOLDS,
    marketing_basis,
    marketing_threshold,
)
from src.exams.models import StudentTestimonial
from src.exams.routes import export_exam_results, list_exam_results
from src.schemas.models import Group, GroupStudent, UserInDB


@pytest.fixture
def db():
    from sqlalchemy import event
    from sqlalchemy.exc import OperationalError
    from sqlalchemy.orm import Session as SASession
    from src.config import engine
    try:
        connection = engine.connect()
    except OperationalError:
        pytest.skip("No database available (requires Postgres); skipping")
    trans = connection.begin()
    session = SASession(bind=connection)
    session.begin_nested()

    @event.listens_for(session, "after_transaction_end")
    def _restart_savepoint(sess, transaction):
        if transaction.nested and not transaction._parent.nested:
            sess.begin_nested()

    try:
        yield session
    finally:
        event.remove(session, "after_transaction_end", _restart_savepoint)
        session.close()
        trans.rollback()
        connection.close()


def _user(db, role, email, name=None):
    u = UserInDB(email=email, name=name or f"{role} {email}", hashed_password="x",
                 role=role, is_active=True)
    db.add(u)
    db.flush()
    return u


def _result(db, student, *, exam_type="sat", total=1450, test_date=None,
            status="reported", superseded=False):
    r = ExamResult(
        student_id=student.id, exam_type=exam_type,
        test_date=test_date or date(2026, 6, 6),
        total_score=Decimal(str(total)),
        verbal_score=700 if exam_type == "sat" else None,
        math_score=750 if exam_type == "sat" else None,
        status=status, source="staff",
        recorded_at=datetime.now(timezone.utc), is_superseded=superseded,
    )
    db.add(r)
    db.flush()
    return r


def _testimonial(db, student, *, revoked=False):
    """An approved, consented testimonial - or one whose consent was withdrawn."""
    now = datetime.now(timezone.utc)
    t = StudentTestimonial(
        student_id=student.id, quote="Best year of my life.",
        status="revoked" if revoked else "approved",
        consent_given=not revoked, consent_channels=["website"], guardian_consent=False,
        consent_recorded_at=now, approved_at=now,
        revoked_at=now if revoked else None,
    )
    db.add(t)
    db.flush()
    return t


@pytest.fixture
def world(db):
    """One curator with one group; ``student(tag)`` enrols a fresh student in it."""
    curator = _user(db, "curator", "mk-cur@t.io")
    group = Group(name="mk Group", curator_id=curator.id,
                  is_active=True, is_over=False, program_type="sat")
    db.add(group)
    db.flush()

    def student(tag):
        s = _user(db, "student", f"mk-{tag}@t.io", name=f"mk Student {tag}")
        db.add(GroupStudent(group_id=group.id, student_id=s.id))
        db.flush()
        return s

    return dict(curator=curator, group=group, student=student)


def _rows(user, db, **kw):
    params = dict(exam_type="sat", group_id=None, date_field="planned",
                  date_from=None, date_to=None, exact_date=None,
                  status=None, search=None, marketing_only=False,
                  limit=200, offset=0, current_user=user, db=db)
    params.update(kw)
    return list_exam_results(**params)


def _row_for(user, db, student, **kw):
    return next(r for r in _rows(user, db, **kw) if r.student.student_id == student.id)


# --------------------------------------------------------------------------------------
# The rule table
# --------------------------------------------------------------------------------------

def test_threshold_is_sat_only_and_strictly_above_1400():
    assert MARKETING_SCORE_THRESHOLDS == {"sat": 1400}
    assert marketing_threshold("sat") == 1400
    assert marketing_threshold("ielts") is None
    assert marketing_threshold("nuet") is None


def test_no_result_and_no_testimonial_is_no_basis():
    assert marketing_basis("sat", None, None) == []


# --------------------------------------------------------------------------------------
# Score basis
# --------------------------------------------------------------------------------------

def test_sat_above_1400_is_eligible_on_score_alone(db, world):
    s = world["student"]("1410")
    _result(db, s, total=1410)

    row = _row_for(world["curator"], db, s)
    assert row.marketing_eligible is True
    assert row.marketing_basis == [BASIS_SCORE]
    assert row.marketing_threshold == 1400


def test_exactly_1400_is_not_eligible(db, world):
    """The owner asked for 'higher than 1400', so 1400 itself does not qualify."""
    s = world["student"]("1400")
    _result(db, s, total=1400)

    row = _row_for(world["curator"], db, s)
    assert row.marketing_eligible is False
    assert row.marketing_basis == []
    assert row.marketing_threshold == 1400   # the rule is still reported


def test_an_unverified_reported_result_still_counts(db, world):
    """The row carries the status, so sales can see it was not checked against proof."""
    s = world["student"]("reported")
    _result(db, s, total=1420, status="reported")

    row = _row_for(world["curator"], db, s)
    assert row.marketing_eligible is True
    assert row.result.status == "reported"


# --------------------------------------------------------------------------------------
# Testimonial basis, and the two together
# --------------------------------------------------------------------------------------

def test_below_threshold_with_a_consented_testimonial_is_eligible_on_testimonial(db, world):
    s = world["student"]("1390")
    _result(db, s, total=1390)
    _testimonial(db, s)

    row = _row_for(world["curator"], db, s)
    assert row.marketing_eligible is True
    assert row.marketing_basis == [BASIS_TESTIMONIAL]


def test_above_threshold_with_a_testimonial_reports_both_bases(db, world):
    """Both are reported so sales can tell a consented testimonial (name and photo
    usable) from a bare score (no consent recorded)."""
    s = world["student"]("1450")
    _result(db, s, total=1450)
    _testimonial(db, s)

    row = _row_for(world["curator"], db, s)
    assert row.marketing_eligible is True
    assert row.marketing_basis == [BASIS_SCORE, BASIS_TESTIMONIAL]


def test_a_revoked_testimonial_gives_no_basis(db, world):
    s = world["student"]("revoked")
    _result(db, s, total=1350)
    _testimonial(db, s, revoked=True)

    row = _row_for(world["curator"], db, s)
    assert row.marketing_eligible is False
    assert row.marketing_basis == []


# --------------------------------------------------------------------------------------
# Exam types without a threshold
# --------------------------------------------------------------------------------------

def test_ielts_has_no_score_threshold(db, world):
    """An 8.5 band is excellent, but IELTS eligibility comes from a testimonial only."""
    s = world["student"]("ielts")
    _result(db, s, exam_type="ielts", total="8.5")

    row = _row_for(world["curator"], db, s, exam_type="ielts")
    assert row.marketing_eligible is False
    assert row.marketing_basis == []
    assert row.marketing_threshold is None


def test_ielts_with_a_testimonial_is_eligible_via_the_testimonial(db, world):
    s = world["student"]("ielts-t")
    _result(db, s, exam_type="ielts", total="7.0")
    _testimonial(db, s)

    row = _row_for(world["curator"], db, s, exam_type="ielts")
    assert row.marketing_eligible is True
    assert row.marketing_basis == [BASIS_TESTIMONIAL]


# --------------------------------------------------------------------------------------
# Only the current attempt counts
# --------------------------------------------------------------------------------------

def test_a_rejected_result_never_qualifies(db, world):
    s = world["student"]("rejected")
    _result(db, s, total=1500, status="rejected")

    row = _row_for(world["curator"], db, s)
    assert row.marketing_eligible is False
    assert row.marketing_basis == []


def test_a_superseded_high_score_does_not_carry_over_to_the_current_attempt(db, world):
    """The 1500 was corrected away; the current attempt is the 1300 that replaced it."""
    s = world["student"]("superseded")
    _result(db, s, total=1500, test_date=date(2026, 3, 14), superseded=True)
    _result(db, s, total=1300, test_date=date(2026, 6, 6))

    row = _row_for(world["curator"], db, s)
    assert row.result.total_score == Decimal("1300")
    assert row.marketing_eligible is False


def test_a_rejected_re_sit_does_not_mask_the_verified_result_it_follows(db, world):
    """Rejecting an attempt does not mark it superseded, so it is still the newest row
    the GRID shows - but the student's current attempt is the verified 1500 before it,
    exactly as ``latest_results_by_student`` reads it. The student must stay eligible."""
    s = world["student"]("rejected-resit")
    _result(db, s, total=1500, test_date=date(2026, 3, 14), status="verified")
    _result(db, s, total=1150, test_date=date(2026, 6, 6), status="rejected")

    row = _row_for(world["curator"], db, s)
    # The grid still displays the rejected re-sit, so the curator can see the rejection.
    assert row.result.total_score == Decimal("1150")
    assert row.result.status == "rejected"
    # ...but marketing is judged on the current attempt, the March 1500.
    assert row.marketing_eligible is True
    assert row.marketing_basis == [BASIS_SCORE]

    ids = {r.student.student_id for r in _rows(world["curator"], db, marketing_only=True)}
    assert s.id in ids


def test_a_date_window_narrowing_the_grid_does_not_qualify_an_old_attempt(db, world):
    """The mirror image: filtering the grid to March shows the 1450 sat in March, but
    the student's current attempt is the June 1300, so they are NOT marketing-eligible."""
    s = world["student"]("windowed")
    _result(db, s, total=1450, test_date=date(2026, 3, 14))
    _result(db, s, total=1300, test_date=date(2026, 6, 6))

    row = _row_for(world["curator"], db, s, date_field="actual",
                   date_from=date(2026, 3, 1), date_to=date(2026, 3, 31))
    assert row.result.total_score == Decimal("1450")
    assert row.marketing_eligible is False
    assert row.marketing_basis == []

    ids = {
        r.student.student_id
        for r in _rows(world["curator"], db, date_field="actual",
                       date_from=date(2026, 3, 1), date_to=date(2026, 3, 31),
                       marketing_only=True)
    }
    assert s.id not in ids


def test_a_status_filter_narrowing_the_grid_does_not_change_eligibility(db, world):
    """A ``status`` filter is a view over the grid, not a redefinition of the current
    attempt: the June 1300 still decides whether this student may be featured."""
    s = world["student"]("status-filtered")
    _result(db, s, total=1500, test_date=date(2026, 3, 14), status="verified")
    _result(db, s, total=1300, test_date=date(2026, 6, 6), status="reported")

    # Unfiltered, the current attempt is the June 1300 - below the threshold.
    assert _row_for(world["curator"], db, s).marketing_eligible is False
    # Filtering to verified rows displays the March 1500, but must not resurrect it as
    # the basis for featuring the student.
    row = _row_for(world["curator"], db, s, status="verified")
    assert row.result.total_score == Decimal("1500")
    assert row.marketing_eligible is False


# --------------------------------------------------------------------------------------
# The marketing_only filter
# --------------------------------------------------------------------------------------

def _seed_mix(db, world):
    """Four students: both bases, score only, testimonial only, neither."""
    both = world["student"]("both")
    _result(db, both, total=1450)
    _testimonial(db, both)

    score = world["student"]("score")
    _result(db, score, total=1410)

    quote = world["student"]("quote")
    _result(db, quote, total=1300)
    _testimonial(db, quote)

    none = world["student"]("none")
    _result(db, none, total=1300)
    return dict(both=both, score=score, quote=quote, none=none)


def test_marketing_only_returns_exactly_the_eligible_rows(db, world):
    people = _seed_mix(db, world)

    ids = {r.student.student_id for r in _rows(world["curator"], db, marketing_only=True)}
    assert ids == {people["both"].id, people["score"].id, people["quote"].id}


def test_without_marketing_only_every_row_is_returned(db, world):
    people = _seed_mix(db, world)

    rows = _rows(world["curator"], db, marketing_only=False)
    assert {r.student.student_id for r in rows} == {p.id for p in people.values()}
    # ...and each row still says whether it is eligible, so the UI can filter too.
    by_id = {r.student.student_id: r for r in rows}
    assert by_id[people["none"].id].marketing_eligible is False
    assert by_id[people["score"].id].marketing_eligible is True


# --------------------------------------------------------------------------------------
# Export
# --------------------------------------------------------------------------------------

def _export(user, db, **kw):
    params = dict(exam_type="sat", group_id=None, date_field="planned",
                  date_from=None, date_to=None, exact_date=None,
                  status=None, search=None, marketing_only=False,
                  current_user=user, db=db)
    params.update(kw)
    resp = export_exam_results(**params)
    assert resp.body[:2] == b"PK"
    return load_workbook(BytesIO(resp.body)).active


def test_export_has_a_marketing_column_with_the_bases_spelled_out(db, world):
    _seed_mix(db, world)

    ws = _export(world["curator"], db)
    headers = [c.value for c in ws[1]]
    assert "Маркетинг" in headers
    col = headers.index("Маркетинг") + 1
    by_name = {
        ws.cell(row=i, column=1).value: (ws.cell(row=i, column=col).value or "")
        for i in range(2, ws.max_row + 1)
    }
    assert by_name["mk Student both"] == "балл > 1400, отзыв"
    assert by_name["mk Student score"] == "балл > 1400"
    assert by_name["mk Student quote"] == "отзыв"
    assert by_name["mk Student none"] == ""


def test_export_honours_marketing_only(db, world):
    _seed_mix(db, world)

    ws = _export(world["curator"], db, marketing_only=True)
    names = {ws.cell(row=i, column=1).value for i in range(2, ws.max_row + 1)}
    assert names == {"mk Student both", "mk Student score", "mk Student quote"}
