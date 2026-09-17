"""The export is the sheet head teachers already know, filled in by the LMS.

Same tabs (one per programme), same headings — ФИО, the dates, «Пропуски» and «Наказание» — so a
payroll clerk opening it recognises it, and an archived period reads like the old spreadsheet.
"""
from io import BytesIO

from openpyxl import load_workbook

from src.discipline.export import workbook_for

REGISTER = {
    "period": {"key": "2026-09-16", "label": "16–30 September 2026",
               "start": "2026-09-16", "end": "2026-09-18", "closed": False},
    "days": ["2026-09-16", "2026-09-17", "2026-09-18"],
    "teachers": [
        {"teacher_id": 1, "name": "Кенжебаев Арсен", "program": "SAT",
         "days": {"2026-09-17": {"late_minutes": 3, "early_minutes": 0, "misses": 0, "fine": 900,
                                 "unpriced": 0, "unmeasurable": 0, "state": "late"}},
         "totals": {"late_minutes": 3, "early_minutes": 0, "misses": 0, "fine": 900, "unpriced": 0,
                    "lessons": 4, "unmeasurable": 0}},
        {"teacher_id": 2, "name": "Бутырин Даниил", "program": "IELTS",
         "days": {"2026-09-18": {"late_minutes": 0, "early_minutes": 0, "misses": 1, "fine": 0,
                                 "unpriced": 1, "unmeasurable": 0, "state": "miss"}},
         "totals": {"late_minutes": 0, "early_minutes": 0, "misses": 1, "fine": 0, "unpriced": 1,
                    "lessons": 2, "unmeasurable": 0}},
    ],
    "totals": {"late_minutes": 3, "early_minutes": 0, "misses": 1, "fine": 900, "unpriced": 1,
               "lessons": 6, "unmeasurable": 0},
}


def _sheets():
    stream = BytesIO()
    workbook_for(REGISTER).save(stream)
    return load_workbook(BytesIO(stream.getvalue()))


def test_one_tab_per_programme_as_the_sheet_has():
    assert _sheets().sheetnames == ["SAT", "IELTS"]


def test_the_headings_are_the_ones_head_teachers_read():
    sheet = _sheets()["SAT"]
    assert sheet["A1"].value == "ФИО"
    assert sheet["B1"].value == "16.09.2026"
    assert (sheet["B2"].value, sheet["C2"].value) == ("Пропуски", "Наказание")


def test_a_teachers_row_carries_the_minutes_and_the_money():
    sheet = _sheets()["SAT"]
    assert sheet["A3"].value == "Кенжебаев Арсен"
    assert sheet["D3"].value == "3 мин"      # 17.09, «Пропуски» says what happened
    assert sheet["E3"].value == 900          # 17.09, «Наказание» in ₸


def test_a_missed_lesson_says_so_and_waits_for_a_price():
    sheet = _sheets()["IELTS"]
    assert sheet["F3"].value == "Пропуск"    # 18.09
    assert sheet["G3"].value is None         # nobody has priced it yet


def test_the_period_and_the_totals_are_written_where_payroll_looks():
    sheet = _sheets()["SAT"]
    assert sheet["A2"].value == "16–30 September 2026"
    header = [cell.value for cell in sheet[1]]
    assert header[-3:] == ["Опозданий, мин", "Пропусков", "Штраф, ₸"]
    assert [cell.value for cell in sheet[3]][-3:] == [3, 0, 900]
    assert [cell.value for cell in sheet[4]][0] == "Итого"


def test_a_name_that_starts_with_an_equals_sign_is_not_a_formula():
    register = {**REGISTER, "teachers": [{**REGISTER["teachers"][0], "name": "=cmd|' /c calc'!A1"}]}
    stream = BytesIO()
    workbook_for(register).save(stream)
    value = load_workbook(BytesIO(stream.getvalue()))["SAT"]["A3"].value
    assert not str(value).startswith("=")


def test_an_empty_period_still_opens():
    register = {**REGISTER, "teachers": [], "totals": {k: 0 for k in REGISTER["totals"]}}
    stream = BytesIO()
    workbook_for(register).save(stream)
    sheet = load_workbook(BytesIO(stream.getvalue())).worksheets[0]
    assert sheet["A2"].value == "16–30 September 2026"
